"""Build CMC-formatted main and supplementary DOCX files from Markdown sources."""

from __future__ import annotations

import re
import shutil
import subprocess
import hashlib
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "artifacts" / "cmc_template" / "CMC-Template.docx"
OUTDIR = ROOT / "paper" / "cmc_submission"
ASSETDIR = OUTDIR / "generated_assets"
HIGHRESDIR = ASSETDIR / "figures_600dpi"
PDFTOPPM = shutil.which("pdftoppm")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_shading(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="BFBFBF", size="4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_run_font(run, name="Times New Roman", size=11, bold=None, italic=None, color="000000") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_paragraph(paragraph, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=3, first=0.22,
                        keep_next=False, line=1.0) -> None:
    paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if first:
        fmt.first_line_indent = Inches(first)
    else:
        fmt.first_line_indent = None
    fmt.keep_with_next = keep_next


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def plain_text(text: str) -> str:
    replacements = {
        "\\times": "x", "\\pm": "+/-", "\\ge": ">=", "\\le": "<=",
        "\\mu": "mu", "\\tau": "tau", "\\mathbb": "", "\\operatorname": "",
        "\\left": "", "\\right": "", "\\max": "max", "\\min": "min",
        "\\sum": "sum", "\\frac": "", "\\exp": "exp", "\\sqrt": "sqrt",
        "{": "", "}": "", "$": "",
    }
    for a, b in replacements.items():
        text = text.replace(a, b)
    text = re.sub(r"\bmu(?=_[A-Za-z0-9]|\()", "μ", text)
    text = re.sub(r"\btau(?=_[A-Za-z0-9]|\()", "τ", text)
    for ascii_suffix, unicode_suffix in {
        "_t": "ₜ", "_i": "ᵢ", "_j": "ⱼ", "_0": "₀", "_1": "₁",
    }.items():
        text = text.replace(ascii_suffix, unicode_suffix)
    return text


def add_inline_markup(paragraph, text: str, *, font="Times New Roman", size=11) -> None:
    text = plain_text(text)
    parts = re.split(r"(\*\*.*?\*\*|\*[^*]+?\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        italic = (part.startswith("*") and part.endswith("*") and not bold)
        value = part[2:-2] if bold else (part[1:-1] if italic else part)
        run = paragraph.add_run(value)
        set_run_font(run, font, size, bold=bold, italic=italic)


def add_body_paragraph(doc, text: str, *, size=11, font="Times New Roman", first=0.22,
                       after=3, italic=False, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    configure_paragraph(p, align=align, after=after, first=first)
    if bold or italic:
        r = p.add_run(plain_text(text))
        set_run_font(r, font, size, bold=bold, italic=italic)
    else:
        add_inline_markup(p, text, font=font, size=size)
    return p


def add_heading(doc, text: str, level: int):
    p = doc.add_paragraph()
    p.style = f"Heading {min(level, 3)}"
    configure_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=12 if level == 1 else 8,
                        after=3, first=0, keep_next=True)
    r = p.add_run(plain_text(text))
    set_run_font(r, "Times New Roman", 11, bold=(level <= 2), italic=(level == 2 or level == 3))
    return p


def render_equation(latex: str, key: str) -> Path:
    ASSETDIR.mkdir(parents=True, exist_ok=True)
    path = ASSETDIR / f"equation_{key}.png"
    if path.exists():
        return path
    exact = {
        "r_t = [y_t - mu_t(x_t)] / s_t,": "rₜ = [yₜ − μₜ(xₜ)] / sₜ,",
        "H_t(x) = mu_t(x) + s_t q_t(x),": "Hₜ(x) = μₜ(x) + sₜqₜ(x),",
    }
    display = exact.get(latex, latex)
    replacements = {
        r"\mu_t": "mu_t", r"\tau_t": "tau_t", r"\qquad": "    ",
        r"\frac": "", r"\max": "max", r"\min": "min", r"\left": "", r"\right": "",
        r"\mathbb": "", r"\Omega": "Omega", r"\in": " in ", r"\|": "|",
    }
    if latex not in exact:
        for a, b in replacements.items():
            display = display.replace(a, b)
        display = display.replace("{", "").replace("}", "").replace("\\", "")
    canvas = Image.new("RGBA", (1860, 180), (255, 255, 255, 0))
    draw = ImageDraw.Draw(canvas)
    font_path = Path(r"C:\Windows\Fonts\cambria.ttc")
    font = ImageFont.truetype(str(font_path), 56) if font_path.exists() else ImageFont.load_default()
    box = draw.textbbox((0, 0), display, font=font)
    x = max(20, (canvas.width - (box[2] - box[0])) // 2)
    y = max(10, (canvas.height - (box[3] - box[1])) // 2 - box[1])
    draw.text((x, y), display, fill=(0, 0, 0, 255), font=font)
    crop = canvas.getbbox()
    if crop:
        canvas = canvas.crop((max(0, crop[0] - 30), max(0, crop[1] - 20), min(canvas.width, crop[2] + 30), min(canvas.height, crop[3] + 20)))
    canvas.save(path, dpi=(300, 300))
    return path


def submission_raster(image_path: Path) -> Path:
    """Rasterize the vector companion at 3900 px (600 dpi at 6.5 in)."""
    pdf_path = image_path.with_suffix(".pdf")
    if not pdf_path.exists():
        return image_path
    HIGHRESDIR.mkdir(parents=True, exist_ok=True)
    token = hashlib.sha1(str(pdf_path).encode("utf-8")).hexdigest()[:10]
    target = HIGHRESDIR / f"{image_path.stem}_{token}.png"
    if not PDFTOPPM:
        raise RuntimeError("pdftoppm was not found on PATH; install Poppler before building DOCX figures")
    subprocess.run([PDFTOPPM, "-png", "-singlefile", "-scale-to-x", "3900",
                    "-scale-to-y", "-1", str(pdf_path), str(target.with_suffix(""))],
                   check=True, capture_output=True)
    with Image.open(target) as im:
        rgb = im.convert("RGB")
        rgb.save(target, dpi=(600, 600), optimize=True)
    return target


def add_picture(doc, image_path: Path, caption: str, number: str) -> None:
    image_path = submission_raster(image_path)
    with Image.open(image_path) as im:
        px_w, px_h = im.size
    width = min(6.3, px_w / 300.0)
    if width < 3.2:
        width = min(5.2, px_w / 240.0)
    p = doc.add_paragraph()
    configure_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=2, first=0, keep_next=True)
    run = p.add_run()
    inline = run.add_picture(str(image_path), width=Inches(width))
    inline._inline.docPr.set("descr", caption)
    cap = doc.add_paragraph()
    configure_paragraph(cap, align=WD_ALIGN_PARAGRAPH.CENTER, after=6, first=0)
    r = cap.add_run(f"Figure {number}: {caption}")
    set_run_font(r, "Times New Roman", 10)


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        vals = [v.strip() for v in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", v or "---") for v in vals):
            continue
        rows.append(vals)
    return rows


def add_table(doc, rows: list[list[str]], caption: str, number: str) -> None:
    cap = doc.add_paragraph()
    configure_paragraph(cap, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=3, first=0, keep_next=True)
    r = cap.add_run(f"Table {number}: {plain_text(caption)}")
    set_run_font(r, "Times New Roman", 10)
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    if rows:
        set_repeat_table_header(table.rows[0])
    for i, row in enumerate(rows):
        tr_pr = table.rows[i]._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        for j in range(cols):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if i == 0:
                set_cell_shading(cell, "1F4E78")
            elif i % 2 == 0:
                set_cell_shading(cell, "EAF2F8")
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(plain_text(row[j]) if j < len(row) else "")
            set_run_font(run, "Times New Roman", 8.5 if cols >= 5 else 9,
                         bold=(i == 0), color="FFFFFF" if i == 0 else "000000")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def remove_instructional_body(doc: Document) -> None:
    body = doc._element.body
    keep = {id(p._p) for p in doc.paragraphs[:9]}
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        if id(child) in keep:
            continue
        body.remove(child)


def replace_slot(paragraph, text: str, *, font="Minion Pro", size=9, bold=False, italic=False,
                 align=WD_ALIGN_PARAGRAPH.LEFT, after=2) -> None:
    clear_paragraph(paragraph)
    configure_paragraph(paragraph, align=align, after=after, first=0)
    run = paragraph.add_run(text)
    set_run_font(run, font, size, bold=bold, italic=italic)


def prepare_template(title: str, supplementary: bool = False) -> Document:
    doc = Document(str(TEMPLATE))
    remove_instructional_body(doc)
    # The distributed template contains an empty section-level bidi flag. It
    # reverses newly authored table columns in LibreOffice despite LTR prose.
    # Remove only that flag so tables and page furniture render left-to-right.
    for section in doc.sections:
        sect_pr = section._sectPr
        bidi = sect_pr.find(qn("w:bidi"))
        if bidi is not None:
            sect_pr.remove(bidi)
    slots = doc.paragraphs[:9]
    replace_slot(slots[0], "DOI: to be assigned", font="Times New Roman", size=7)
    replace_slot(slots[1], "SUPPLEMENTARY MATERIALS" if supplementary else "ARTICLE",
                 font="Times New Roman", size=9, bold=True, after=5)
    replace_slot(slots[2], title, font="Minion Pro", size=14, bold=True, after=6)
    clear_paragraph(slots[3])
    configure_paragraph(slots[3], align=WD_ALIGN_PARAGRAPH.LEFT, after=4, first=0)
    for author, marker, separator in (
        ("Shu-Chuan Chu", "1,*", ", "),
        ("Yayu Zhang", "1", ", "),
        ("Zhenyu Meng", "2", " and "),
        ("Jia Yue", "3", ""),
    ):
        run = slots[3].add_run(author)
        set_run_font(run, "Minion Pro", 11, bold=True)
        run = slots[3].add_run(marker)
        set_run_font(run, "Minion Pro", 8, bold=True)
        run.font.superscript = True
        if separator:
            run = slots[3].add_run(separator)
            set_run_font(run, "Minion Pro", 11, bold=True)
    replace_slot(slots[4], "1 Nanjing University of Information Science and Technology, Nanjing, China",
                 font="Minion Pro", size=9)
    replace_slot(slots[5], "2 Affiliation to be confirmed for Zhenyu Meng; 3 Affiliation to be confirmed for Jia Yue",
                 font="Minion Pro", size=9)
    replace_slot(slots[6], "* Corresponding Author: Shu-Chuan Chu. Email: to be confirmed",
                 font="Minion Pro", size=9, after=2)
    replace_slot(slots[7], "Affiliations 2-3 and the corresponding-author email must be confirmed before submission.",
                 font="Minion Pro", size=8, italic=True, after=2)
    replace_slot(slots[8], "Received: Day Month Year    Accepted: Day Month Year",
                 font="Minion Pro", size=8, after=8)
    return doc


def add_abstract(doc: Document, abstract: str, keywords: str) -> None:
    for label, body in (("ABSTRACT", abstract), ("KEYWORDS", keywords)):
        p = doc.add_paragraph()
        set_paragraph_shading(p, "E7E6E6")
        configure_paragraph(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=0 if label == "ABSTRACT" else 8,
                            first=0)
        r = p.add_run(label + "\n")
        set_run_font(r, "Minion Pro", 10, bold=True)
        add_inline_markup(p, body, font="Minion Pro", size=10)


def build_from_markdown(source: Path, target: Path, *, supplementary=False) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    title = lines[0].lstrip("# ").strip()
    doc = prepare_template(title, supplementary=supplementary)
    start = 1
    if not supplementary:
        a_idx = lines.index("## Abstract")
        k_idx = next(i for i, line in enumerate(lines) if line.startswith("**Keywords:**"))
        abstract = " ".join(line.strip() for line in lines[a_idx + 1:k_idx] if line.strip())
        keywords = lines[k_idx].replace("**Keywords:**", "").strip()
        add_abstract(doc, abstract, keywords)
        start = k_idx + 1
    else:
        # Skip the author line and add the opening scope paragraph before S1.
        first_heading = next(i for i, line in enumerate(lines) if line.startswith("# S0"))
        for line in lines[1:first_heading]:
            if line.strip() and not line.startswith("**"):
                add_body_paragraph(doc, line.strip(), first=0)
        start = first_heading

    figure_no = 0
    table_no = 0
    equation_no = 0
    pending_caption = None
    current_heading = "Supplementary results"
    paragraph_buffer: list[str] = []

    def flush_buffer():
        nonlocal paragraph_buffer
        if paragraph_buffer:
            text = " ".join(s.strip() for s in paragraph_buffer).strip()
            if text:
                if text.startswith("**Proposition") or text.startswith("**Corollary") or text.startswith("**Proof"):
                    p = add_body_paragraph(doc, text, first=0)
                    p.paragraph_format.keep_with_next = False
                else:
                    add_body_paragraph(doc, text)
            paragraph_buffer = []

    i = start
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            flush_buffer()
            i += 1
            continue
        if line.startswith("# "):
            flush_buffer()
            current_heading = line[2:].strip()
            add_heading(doc, current_heading, 1)
            i += 1
            continue
        if line.startswith("## "):
            flush_buffer()
            heading = line[3:].strip()
            if heading in {"Abstract"}:
                i += 1
                continue
            current_heading = heading
            add_heading(doc, heading, 1 if re.match(r"(?:S?\d+|Acknowledgement|Funding|Author|Availability|Ethics|Conflicts|Supplementary|References)", heading) else 2)
            i += 1
            continue
        if line.startswith("### "):
            flush_buffer()
            current_heading = line[4:].strip()
            add_heading(doc, current_heading, 2)
            i += 1
            continue
        if line.startswith("$$") and line.endswith("$$"):
            flush_buffer()
            equation_no += 1
            latex = line[2:-2].strip()
            eq_path = render_equation(latex, f"v3_{'s' if supplementary else 'm'}_{equation_no}")
            p = doc.add_paragraph()
            configure_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=3, after=3, first=0)
            inline = p.add_run().add_picture(str(eq_path), width=Inches(4.8))
            inline._inline.docPr.set("descr", f"Equation {equation_no}: {plain_text(latex)}")
            i += 1
            continue
        image_match = re.match(r"!\[(.*?)\]\((.*?)\)", line)
        if image_match:
            flush_buffer()
            figure_no += 1
            caption, rel = image_match.groups()
            image_path = (source.parent / rel).resolve()
            add_picture(doc, image_path, plain_text(caption), f"S{figure_no}" if supplementary else str(figure_no))
            i += 1
            continue
        if line.startswith("**Table ") or line.startswith("Table:"):
            flush_buffer()
            pending_caption = line.strip("*")
            pending_caption = re.sub(r"^Table(?:\s+\d+)?[.:]\s*", "", pending_caption)
            i += 1
            continue
        if line.startswith("|"):
            flush_buffer()
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            rows = parse_table(block)
            if rows:
                table_no += 1
                default_caption = f"Results for {re.sub(r'^S?\d+(?:\.\d+)*\s*', '', current_heading)}"
                add_table(doc, rows, pending_caption or default_caption, f"S{table_no}" if supplementary else str(table_no))
            pending_caption = None
            continue
        if re.match(r"^\d+\.\s", line) and supplementary:
            flush_buffer()
            p = add_body_paragraph(doc, line, first=0)
            p.paragraph_format.left_indent = Inches(0.2)
            i += 1
            continue
        paragraph_buffer.append(line)
        i += 1
    flush_buffer()

    # Global section and default settings retained from CMC template.
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    # Preserve the template's branded headers and footers while making their
    # decorative images and layout tables accessible to Word and PDF readers.
    seen_parts = set()
    for section in doc.sections:
        containers = (
            section.header, section.first_page_header, section.even_page_header,
            section.footer, section.first_page_footer, section.even_page_footer,
        )
        for container in containers:
            part_key = str(container.part.partname)
            if part_key in seen_parts:
                continue
            seen_parts.add(part_key)
            for node in container._element.iter():
                if node.tag == qn("wp:docPr") and not node.get("descr"):
                    node.set("descr", "CMC journal branding")
            for table in container.tables:
                if table.rows:
                    set_repeat_table_header(table.rows[0])

    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(target))


def main() -> None:
    OUTDIR.mkdir(parents=True, exist_ok=True)
    build_from_markdown(ROOT / "paper" / "cmc_submission_manuscript.md",
                        OUTDIR / "CMC_submission_manuscript.docx")
    build_from_markdown(ROOT / "paper" / "cmc_supplementary_materials.md",
                        OUTDIR / "CMC_supplementary_materials.docx", supplementary=True)


if __name__ == "__main__":
    main()
