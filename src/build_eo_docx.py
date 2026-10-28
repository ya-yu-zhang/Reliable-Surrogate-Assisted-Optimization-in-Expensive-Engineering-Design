"""Build unbranded Engineering Optimization submission DOCX files."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import build_cmc_docx as base


ROOT = Path(__file__).resolve().parents[1]
OUTDIR = ROOT / "paper" / "eo_submission"


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))
    base.set_run_font(run, "Times New Roman", 9)


def prepare_eo_document(title: str, supplementary: bool = False) -> Document:
    doc = Document()
    doc.core_properties.title = title
    doc.core_properties.subject = (
        "Supplementary materials for Engineering Optimization"
        if supplementary
        else "Engineering Optimization submission manuscript"
    )
    doc.core_properties.author = (
        "Shu-Chuan Chu; Ya-Yu Zhang; Zhenyu Meng; Jia Zhao; Jeng-Shyang Pan"
    )
    doc.core_properties.last_modified_by = "Ya-Yu Zhang"
    doc.core_properties.comments = ""
    doc.core_properties.keywords = (
        "expensive engineering optimization; surrogate-assisted optimization; "
        "model management; decision-consequence auditing"
    )

    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    add_page_number(section.footer.paragraphs[0])

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")

    title_style = doc.styles["Title"]
    title_style.font.name = "Times New Roman"
    title_style.font.size = Pt(16)
    title_style.font.bold = True
    title_ppr = title_style._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    article = doc.add_paragraph()
    base.configure_paragraph(
        article, align=WD_ALIGN_PARAGRAPH.CENTER, after=6, first=0
    )
    run = article.add_run("SUPPLEMENTARY MATERIALS" if supplementary else "ARTICLE")
    base.set_run_font(run, "Times New Roman", 10, bold=True)

    title_paragraph = doc.add_paragraph(style="Title")
    base.configure_paragraph(
        title_paragraph, align=WD_ALIGN_PARAGRAPH.CENTER, after=10, first=0
    )
    title_direct_border = title_paragraph._p.get_or_add_pPr().find(qn("w:pBdr"))
    if title_direct_border is not None:
        title_paragraph._p.get_or_add_pPr().remove(title_direct_border)
    title_run = title_paragraph.add_run(title)
    base.set_run_font(title_run, "Times New Roman", 16, bold=True)

    authors = doc.add_paragraph()
    base.configure_paragraph(
        authors, align=WD_ALIGN_PARAGRAPH.CENTER, after=5, first=0
    )
    author_run = authors.add_run(
        "Shu-Chuan Chu¹˒², Ya-Yu Zhang¹, Zhenyu Meng³, Jia Zhao⁴, "
        "and Jeng-Shyang Pan¹˒²˒*"
    )
    base.set_run_font(author_run, "Times New Roman", 11, bold=True)

    affiliations = [
        "¹ School of Computer Science, Nanjing University of Information Science and Technology, Nanjing 210044, Jiangsu, China",
        "² School of Artificial Intelligence/School of Future Technology, Nanjing University of Information Science and Technology, Nanjing 210044, Jiangsu, China",
        "³ School of Computer Science and Software, Zhaoqing University, Zhaoqing 526061, China",
        "⁴ School of Information Engineering, Jiangxi University of Water Resources and Electric Power, Nanchang 330099, China",
        "* Corresponding author: Jeng-Shyang Pan, 100112@nuist.edu.cn",
        "Shu-Chuan Chu, 100130@nuist.edu.cn",
        "Ya-Yu Zhang, 202412492223@nuist.edu.cn",
        "Jia Zhao, zhaojia@juwp.edu.cn",
    ]
    for affiliation in affiliations:
        paragraph = doc.add_paragraph()
        base.configure_paragraph(
            paragraph, align=WD_ALIGN_PARAGRAPH.CENTER, after=1, first=0
        )
        affiliation_run = paragraph.add_run(affiliation)
        base.set_run_font(affiliation_run, "Times New Roman", 9)

    contact = doc.add_paragraph()
    base.configure_paragraph(
        contact, align=WD_ALIGN_PARAGRAPH.CENTER, after=8, first=0
    )
    contact_run = contact.add_run(
        "ORCID: S. Chu, 0000-0003-2117-0618; Y.-Y. Zhang, "
        "0009-0001-2272-4654; Z. Meng, 0000-0002-1466-8082; "
        "Jia Zhao, 0000-0002-3652-1903; J.-S. Pan, 0000-0002-3128-9025"
    )
    base.set_run_font(contact_run, "Times New Roman", 8.5, italic=True)
    return doc


def add_plain_abstract(doc: Document, abstract: str, keywords: str) -> None:
    heading = doc.add_paragraph()
    base.configure_paragraph(
        heading, align=WD_ALIGN_PARAGRAPH.LEFT, before=6, after=3, first=0,
        keep_next=True,
    )
    run = heading.add_run("Abstract")
    base.set_run_font(run, "Times New Roman", 12, bold=True)
    base.add_body_paragraph(doc, abstract, first=0, after=5)

    paragraph = doc.add_paragraph()
    base.configure_paragraph(
        paragraph, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=8, first=0
    )
    label = paragraph.add_run("Keywords: ")
    base.set_run_font(label, "Times New Roman", 10, bold=True)
    base.add_inline_markup(paragraph, keywords, font="Times New Roman", size=10)


def main() -> None:
    OUTDIR.mkdir(parents=True, exist_ok=True)
    base.OUTDIR = OUTDIR
    base.ASSETDIR = OUTDIR / "generated_assets"
    base.HIGHRESDIR = base.ASSETDIR / "figures_600dpi"
    base.prepare_template = prepare_eo_document
    base.add_abstract = add_plain_abstract

    base.build_from_markdown(
        ROOT / "paper" / "engineering_optimization_manuscript.md",
        OUTDIR / "EO_submission_manuscript.docx",
    )
    base.build_from_markdown(
        ROOT / "paper" / "engineering_optimization_supplementary.md",
        OUTDIR / "EO_supplementary_materials.docx",
        supplementary=True,
    )


if __name__ == "__main__":
    main()
