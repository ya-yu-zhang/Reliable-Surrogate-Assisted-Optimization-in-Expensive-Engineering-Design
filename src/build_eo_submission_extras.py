"""Build the Engineering Optimization cover letter and highlights."""

from __future__ import annotations

import shutil
import stat
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTDIR = ROOT / "paper" / "eo_submission"
TMPDIR = ROOT / "tmp" / "eo_cover_highlights"
COVER_TEMPLATE = Path(r"D:\google\GENO Cover Letter Template.docx")

SCOPE_FIT = (
    "The manuscript, \"Decision-consequence auditing for surrogate-assisted "
    "optimization in expensive engineering design,\" concerns numerical optimization "
    "when each evaluation is an expensive engineering simulation and a surrogate "
    "decides which design is evaluated next. This matches the journal's scope in "
    "optimization methods and in optimization applied to engineering design. The "
    "paper specifies the evidence required before a surrogate-management rule can "
    "be treated as an improvement in a finite-element design workflow, and it "
    "applies that test to a density-based beam problem under an explicit evaluation "
    "budget. The object is the paid design decision, covering calibration, candidate "
    "selection, closed-loop performance, and the full cost of the physical calls."
)

CONTRIBUTION = (
    "The manuscript contributes a decision-consequence audit for surrogate-assisted "
    "optimization in expensive engineering design. Exact decision equivalence sets "
    "the minimum query count for reporting every immediate utility, so a frozen "
    "120-state intervention is reported from 138 finite-element solves instead of "
    "480. The procedure separates calibration, structural decision influence, "
    "common-state utility, closed-loop performance and total physical-call cost. "
    "Applied to three recorded acquisition rules, it returns one promotion, one "
    "stop and one held-for-revision verdict. A finite-element beam study and twenty "
    "transfer cases at two load positions show where a calibrated signal fails to "
    "improve the design and where cheaper selection loses endpoint performance."
)

# Affiliation and email follow the manuscript. Homepage and position are filled
# only where a public institutional page states them. Unknown cells stay empty.
AUTHORS = [
    (
        "Shu-Chuan Chu",
        "School of Computer Science; School of Artificial Intelligence/School of Future Technology, Nanjing University of Information Science and Technology",
        "",
        "https://faculty.nuist.edu.cn/zhushujuan/en/index.htm",
        "Professor",
    ),
    (
        "Ya-Yu Zhang",
        "School of Computer Science, Nanjing University of Information Science and Technology",
        "202412492223@nuist.edu.cn",
        "",
        "",
    ),
    (
        "Zhenyu Meng",
        "School of Computer Science and Software, Zhaoqing University",
        "",
        "",
        "Professor",
    ),
    (
        "Jia Zhao",
        "School of Information Engineering, Jiangxi University of Water Resources and Electric Power",
        "zhaojia@juwp.edu.cn",
        "",
        "Professor",
    ),
    (
        "Jeng-Shyang Pan",
        "School of Computer Science; School of Artificial Intelligence/School of Future Technology, Nanjing University of Information Science and Technology",
        "100112@nuist.edu.cn",
        "https://faculty.nuist.edu.cn/panzhengxiang/en/index.htm",
        "Professor; corresponding author",
    ),
]


def set_font(run, size: float = 12, *, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for name in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{name}"), "Times New Roman")


def set_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        paragraph.runs[0].italic = False
        paragraph.runs[0].bold = False
        for run in paragraph.runs[1:]:
            run._element.getparent().remove(run._element)
        return
    paragraph.add_run(text)


def build_cover_letter() -> Path:
    if not COVER_TEMPLATE.is_file():
        raise FileNotFoundError(COVER_TEMPLATE)
    words = CONTRIBUTION.split()
    if len(words) > 100:
        raise RuntimeError(f"Contribution statement is {len(words)} words; the template allows 100")
    OUTDIR.mkdir(parents=True, exist_ok=True)
    target = OUTDIR / "EO_cover_letter.docx"
    if target.exists():
        target.chmod(stat.S_IWRITE | stat.S_IREAD)
    shutil.copy2(COVER_TEMPLATE, target)
    target.chmod(stat.S_IWRITE | stat.S_IREAD)
    doc = Document(target)
    if len(doc.tables) != 3:
        raise RuntimeError(f"Expected 3 template tables, found {len(doc.tables)}")
    set_paragraph_text(doc.tables[0].rows[1].cells[0].paragraphs[0], SCOPE_FIT)
    set_paragraph_text(doc.tables[1].rows[1].cells[0].paragraphs[0], CONTRIBUTION)
    author_table = doc.tables[2]
    if len(author_table.rows) != 7:
        raise RuntimeError(f"Expected a header and six author rows, found {len(author_table.rows)}")
    for row, values in zip(author_table.rows[1:6], AUTHORS, strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            set_paragraph_text(cell.paragraphs[0], value)
    unused = author_table.rows[6]._tr
    unused.getparent().remove(unused)
    # All-author approval (recorded 2026-10-21); kept outside the 100-word contribution cell.
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph
    approval = (
        "All authors reviewed the results and approved the final version of the "
        "manuscript for submission (2026-10-21)."
    )
    av = None
    for para in doc.paragraphs:
        if para.text.strip() == "Author verification":
            av = para
            break
    if av is not None and approval not in "\n".join(x.text for x in doc.paragraphs):
        new_el = OxmlElement("w:p")
        av._p.addprevious(new_el)
        new_para = Paragraph(new_el, av._parent)
        run = new_para.add_run(approval)
        set_font(run, size=12)
        spacer = OxmlElement("w:p")
        new_el.addnext(spacer)

    doc.core_properties.title = "Cover Letter for Engineering Optimization"
    doc.core_properties.author = "Jeng-Shyang Pan"
    doc.core_properties.last_modified_by = "Jeng-Shyang Pan"
    doc.core_properties.subject = (
        "Decision-consequence auditing for surrogate-assisted optimization "
        "in expensive engineering design"
    )
    doc.save(target)
    return target


def build_highlights_docx() -> Path:
    TMPDIR.mkdir(parents=True, exist_ok=True)
    target = TMPDIR / "Highlights.docx"
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(14)
    run = title.add_run("Highlights")
    set_font(run, size=14, bold=True)

    highlights = [
        "Common additive correction leaves exploitation ranks unchanged.",
        "Prediction-before-label records keep a label outside the optimizer.",
        "A shared start separates the search rule from its origin.",
        "Later improvement follows an action with zero immediate gain.",
        "Recovery and verification calls stay inside the paid ledger.",
    ]
    for item in highlights:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.15
        run = paragraph.add_run(item)
        set_font(run)
    doc.core_properties.title = "Highlights"
    doc.core_properties.author = "Jeng-Shyang Pan"
    doc.core_properties.last_modified_by = "Jeng-Shyang Pan"
    doc.save(target)
    return target


def main() -> None:
    print(build_cover_letter())
    print(build_highlights_docx())


if __name__ == "__main__":
    main()
