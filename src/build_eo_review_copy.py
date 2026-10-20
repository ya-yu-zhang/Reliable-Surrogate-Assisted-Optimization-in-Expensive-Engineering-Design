"""Build the double-spaced, line-numbered review copy and the separate figure set.

Taylor & Francis journals usually want a review copy with wide spacing and line
numbers, and figures supplied as separate files. Both are produced from the same
Markdown/DOCX submission track, so nothing here changes the scientific content.
"""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SUBMISSION = ROOT / "paper" / "eo_submission"
FIGURES = SUBMISSION / "figures"

MAIN = SUBMISSION / "EO_submission_manuscript.docx"
REVIEW = SUBMISSION / "EO_submission_manuscript_review.docx"

FIGURE_FILES = [
    ("Fig1", ROOT / "results" / "framework_audit_v1" / "decision_consequence_framework.png"),
    ("Fig2", ROOT / "results" / "mechanism_validation_report_v1" / "endpoint_function_effects.png"),
    ("Fig3", ROOT / "results" / "mechanism_validation_report_v1" / "history_object_differences.png"),
    ("Fig4", ROOT / "results" / "engineering_beam_drso_shared_init_v1" / "complete_shared_initialization_convergence.png"),
    ("Fig5", ROOT / "results" / "engineering_beam_drso_shared_init_v1" / "complete_representative_designs.png"),
    ("Fig6", ROOT / "paper" / "engineering_optimization" / "figures" / "transfer_cost_performance.png"),
]


def add_line_numbers(section) -> None:
    sect_pr = section._sectPr
    for existing in sect_pr.findall(qn("w:lnNumType")):
        sect_pr.remove(existing)
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    sect_pr.append(ln)


def add_page_number(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.text = ""
    run = paragraph.add_run()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)


def build_review_copy() -> Path:
    shutil.copy2(MAIN, REVIEW)
    doc = Document(REVIEW)
    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        paragraph.paragraph_format.space_before = None
        paragraph.paragraph_format.space_after = None
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    for section in doc.sections:
        add_line_numbers(section)
        add_page_number(section)
    doc.save(REVIEW)
    return REVIEW


def export_figures() -> list[tuple[str, int, int, float]]:
    FIGURES.mkdir(parents=True, exist_ok=True)
    rows = []
    for name, source in FIGURE_FILES:
        if not source.exists():
            raise FileNotFoundError(source)
        target = FIGURES / f"{name}.png"
        shutil.copy2(source, target)
        with Image.open(target) as image:
            width, height = image.size
            dpi = image.info.get("dpi", (0, 0))[0]
        rows.append((target.name, width, height, round(float(dpi))))
    lines = ["# Figure files for submission", "",
             "Each file is the exact image embedded in the manuscript; no figure was redrawn.",
             "Effective resolution is reported at a 190 mm full width.", "",
             "| File | Pixels | Stored dpi | Effective dpi at 190 mm |",
             "|---|---:|---:|---:|"]
    for name, width, height, dpi in rows:
        effective = round(width / (190 / 25.4))
        lines.append(f"| {name} | {width}x{height} | {dpi} | {effective} |")
    (FIGURES / "README.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    return rows


def main() -> None:
    review = build_review_copy()
    rows = export_figures()
    print(f"review copy: {review.name}")
    print(f"figures: {len(rows)} files in {FIGURES.relative_to(ROOT)}")
    for name, width, height, dpi in rows:
        print(f"  {name} {width}x{height} {dpi} dpi")


if __name__ == "__main__":
    main()
