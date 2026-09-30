from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

from docx import Document


ROOT = Path(r"E:\论文3")
OUT = ROOT / "paper" / "cmc_submission"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def audit_docx(path: Path, render_dir: Path) -> dict:
    doc = Document(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    text = "\n".join(paragraphs)
    placeholders = sorted(set(re.findall(r"\[[A-Z][^\]\n]{2,80}\]", text)))
    raw_markup = sorted(set(re.findall(r"(?:\*\*|!\[|\$\$|^#{1,3}\s)", text, re.M)))
    tables = [p for p in paragraphs if re.match(r"^Table\s+(?:S)?\d+:", p)]
    figures = [p for p in paragraphs if re.match(r"^Figure\s+(?:S)?\d+:", p)]
    return {
        "path": str(path),
        "bytes": path.stat().st_size,
        "sha256": sha256(path),
        "page_count": len(list(render_dir.glob("page-*.png"))),
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "table_caption_count": len(tables),
        "figure_caption_count": len(figures),
        "word_count_approx": len(re.findall(r"\b[\w'-]+\b", text)),
        "placeholders": placeholders,
        "raw_markdown_tokens": raw_markup,
        "duplicate_table_captions": sorted({x for x in tables if tables.count(x) > 1}),
        "duplicate_figure_captions": sorted({x for x in figures if figures.count(x) > 1}),
    }


def main() -> None:
    main_md = (ROOT / "paper" / "cmc_submission_manuscript.md").read_text(encoding="utf-8")
    abstract_match = re.search(r"## Abstract\s+(.*?)\s+\*\*Keywords:", main_md, re.S)
    abstract = abstract_match.group(1) if abstract_match else ""
    reference_numbers = [int(n) for n in re.findall(r"^\[(\d+)\]", main_md, re.M)]
    body = main_md.split("# References", 1)[0]
    cited_numbers = sorted({int(n) for group in re.findall(r"\[([\d,\-\s]+)\]", body) for n in re.findall(r"\d+", group)})
    status = {
        "main": audit_docx(OUT / "CMC_submission_manuscript.docx", OUT / "qa_declarations_20260930"),
        "supplement": audit_docx(OUT / "CMC_supplementary_materials.docx", OUT / "qa_supp_declarations_20260930"),
        "manuscript_source": {
            "abstract_word_count": len(re.findall(r"\b[\w'-]+\b", abstract)),
            "reference_numbers": reference_numbers,
            "cited_reference_numbers": cited_numbers,
            "missing_reference_entries": sorted(set(cited_numbers) - set(reference_numbers)),
            "uncited_reference_entries": sorted(set(reference_numbers) - set(cited_numbers)),
            "contiguous_reference_list": reference_numbers == list(range(1, len(reference_numbers) + 1)),
        },
        "known_submission_placeholders": [
            "affiliations for Zhenyu Meng and Jia Yue",
        ],
    }
    target = OUT / "CMC_submission_status.json"
    target.write_text(json.dumps(status, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(status, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
